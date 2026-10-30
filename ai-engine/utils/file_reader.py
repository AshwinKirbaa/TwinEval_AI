import importlib.util
import os
import re
import tempfile
import time
from pathlib import Path

import fitz
from docx import Document


def _safe_remove(path: str, retries: int = 5, delay: float = 0.2):
    """
    Best-effort temp file cleanup. On Windows the file can still be
    briefly locked right after Ollama reads it (antivirus scanning,
    delayed handle release, etc). Retry a few times, and if it still
    can't be removed, just leave it for the OS temp-dir cleanup rather
    than failing the whole OCR extraction over a leftover temp file.
    """
    for attempt in range(retries):
        try:
            if os.path.exists(path):
                os.remove(path)
            return
        except Exception:
            if attempt == retries - 1:
                return
            time.sleep(delay)


def _is_scanned_page(page, coverage_threshold: float = 0.6) -> bool:
    """
    Detect pages that are essentially one big embedded photo/scan
    (typical of phone-scan apps like CamScanner/Adobe Scan/Notes
    scanner) rather than a digitally-typed page. Such apps often embed
    a low-quality sidecar OCR text layer for searchability -- that
    text layer must NOT be trusted, since it's frequently garbled and
    would otherwise cause us to skip real OCR on a handwritten page.
    """
    try:
        page_rect = page.rect
        page_area = page_rect.width * page_rect.height
        if page_area <= 0:
            return False

        images = page.get_images(full=True)
        if not images:
            return False

        covered_area = 0.0
        for img in images:
            xref = img[0]
            try:
                rects = page.get_image_rects(xref)
            except Exception:
                rects = []
            for r in rects:
                covered_area += max(0.0, r.width) * max(0.0, r.height)

        return (covered_area / page_area) >= coverage_threshold
    except Exception:
        return False


_COMMON_ENGLISH_WORDS = {
    "the", "and", "is", "of", "to", "in", "you", "are", "this", "for",
    "with", "that", "it", "as", "on", "be", "from", "or", "an", "at",
    "by", "we", "can", "will", "your", "my", "how", "what", "name",
    "hello", "using", "user", "input", "output", "program", "class",
    "public", "string", "if", "else", "system", "print", "chat", "bot",
}


def _looks_like_garbled_text(text: str) -> bool:
    """
    Heuristic check for embedded PDF text layers that are actually
    garbage (bad font encoding / low-quality sidecar OCR from a
    scanning app) rather than real, readable text. Real English text
    -- even technical/code text -- reliably contains a good sprinkling
    of very common short words; garbled glyph soup essentially never
    does.
    """
    words = re.findall(r"[A-Za-z]{2,}", text)
    if len(words) < 8:
        # Too short a sample to judge confidently either way; don't
        # force OCR on something ambiguous like a short typed heading.
        return False
    common_hits = sum(1 for w in words if w.lower() in _COMMON_ENGLISH_WORDS)
    return (common_hits / len(words)) < 0.03


def read_pdf(file, progress_callback=None, ocr_model: str = "gemma3:4b"):
    """
    Extract text from a PDF, page by page.

    For each page: use the embedded text layer only if the page is NOT
    a scanned/photographed page AND the text actually looks like real
    English (not a garbled sidecar OCR layer from a scanning app).
    Otherwise fall back to OCR via Ollama, per page. This means a PDF
    that mixes typed and handwritten/scanned pages -- or that has a
    low-quality embedded text layer from a phone scanning app -- still
    gets every non-genuine page OCR'd correctly.

    progress_callback(current_page: int, total_pages: int, status: str)
    is called before/after each page is processed so a caller (e.g. a
    Streamlit UI) can show live progress on multi-page documents.
    """
    try:
        file.seek(0)
    except Exception:
        pass

    pdf = fitz.open(stream=file.read(), filetype="pdf")
    total_pages = len(pdf)
    pages = []

    # Minimum number of non-whitespace characters for a page's embedded
    # text layer to be trusted (guards against stray watermarks/headers
    # on an otherwise scanned/handwritten page).
    MIN_TEXT_CHARS = 20

    for page_number in range(total_pages):
        page = pdf.load_page(page_number)

        scanned = _is_scanned_page(page)
        page_text = "" if scanned else (page.get_text() or "").strip()
        trust_text_layer = (
            not scanned
            and len(page_text) >= MIN_TEXT_CHARS
            and not _looks_like_garbled_text(page_text)
        )

        if trust_text_layer:
            if progress_callback:
                progress_callback(page_number + 1, total_pages, "text")
            pages.append(page_text)
            continue

        # No usable/trustworthy text layer on this page -> OCR it individually.
        if progress_callback:
            progress_callback(page_number + 1, total_pages, "ocr_start")

        pix = page.get_pixmap(dpi=300)
        temp_fd, temp_img_path = tempfile.mkstemp(suffix=".png")
        os.close(temp_fd)  # release our handle immediately so pix.save() has exclusive access on Windows
        pix.save(temp_img_path)

        try:
            ocr_text = extract_text_from_image_path(temp_img_path, model=ocr_model)
            if ocr_text and ocr_text.strip():
                pages.append(ocr_text.strip())
        finally:
            _safe_remove(temp_img_path)

        if progress_callback:
            progress_callback(page_number + 1, total_pages, "ocr_done")

    return "\n\n".join(pages)


def read_docx(file):
    try:
        file.seek(0)
    except Exception:
        pass
    doc = Document(file)
    lines = [para.text for para in doc.paragraphs if para.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                lines.append(row_text)

    return "\n".join(lines)


def read_txt(file):
    try:
        file.seek(0)
    except Exception:
        pass
    content = file.read()
    if isinstance(content, bytes):
        return content.decode("utf-8", errors="ignore")
    return str(content)


def extract_text_from_image_path(image_path: str, model: str = "gemma3:4b") -> str:
    from agents.ollama_vision_agent import OllamaVisionAgent

    errors = []

    # Check readiness lazily, right here, instead of gating the whole
    # upload up front (see extract_text()/read_pdf()). This function is
    # only ever called for a page/image that actually needs OCR, so a
    # typed PDF with a normal embedded text layer never has to touch
    # Ollama at all -- only pages that truly need OCR pay this cost,
    # and only they get blocked if Ollama isn't ready.
    ok, ready_msg = check_ollama_ready(model=model)
    if not ok:
        errors.append(ready_msg)
    else:
        try:
            text = OllamaVisionAgent(model=model).extract_text(image_path)
            if text and text.strip():
                return text
            errors.append("Ollama returned no text")
        except Exception as e:
            errors.append(f"Ollama OCR Error: {e}")

    if os.environ.get("AI_PROVIDER", "ollama").strip().lower() == "gemini":
        try:
            from utils.gemini_ocr import extract_handwritten_text

            text = extract_handwritten_text(image_path)
            if text and text.strip():
                return text
            errors.append("Gemini returned no text")
        except Exception as gemini_err:
            errors.append(f"Gemini OCR Error: {gemini_err}")

    try:
        return paddleocr_extract_text(image_path)
    except Exception as paddle_err:
        errors.append(f"PaddleOCR Error: {paddle_err}")

    raise RuntimeError("OCR extraction failed: " + " | ".join(errors))


def paddleocr_extract_text(image_path: str) -> str:
    if importlib.util.find_spec("paddleocr") is None:
        raise RuntimeError("PaddleOCR is not installed")

    from paddleocr import PaddleOCR

    ocr = PaddleOCR(use_angle_cls=False, lang="en")
    result = ocr.ocr(image_path, cls=False, det=True)
    lines = []
    for row in result:
        if not row or len(row) < 2:
            continue
        text_segments = []
        for segment in row:
            if len(segment) >= 2 and isinstance(segment[1], (list, tuple)):
                text_segments.append(str(segment[1][0]))
        if text_segments:
            lines.append(" ".join(text_segments))
    if lines:
        return "\n".join(lines)
    raise RuntimeError("PaddleOCR returned no text")


def read_image(file, progress_callback=None, ocr_model: str = "gemma3:4b"):
    """
    Save uploaded image temporarily and extract text using the available OCR pipeline.
    """
    try:
        file.seek(0)
    except Exception:
        pass

    suffix = Path(file.name).suffix or ".png"
    data = file.getbuffer() if hasattr(file, "getbuffer") else file.read()

    temp_fd, temp_path = tempfile.mkstemp(suffix=suffix)
    with os.fdopen(temp_fd, "wb") as f:
        f.write(data)

    try:
        if progress_callback:
            progress_callback(1, 1, "ocr_start")
        text = extract_text_from_image_path(temp_path, model=ocr_model)
        if progress_callback:
            progress_callback(1, 1, "ocr_done")
        return text
    finally:
        _safe_remove(temp_path)


def check_ollama_ready(model: str = "gemma3:4b"):
    """
    Verify the local Ollama server is reachable and the requested model is
    pulled. Returns (ok: bool, message: str) so callers can surface a
    friendly, actionable error instead of a raw stack trace.
    """
    if os.environ.get("AI_PROVIDER", "ollama").strip().lower() == "gemini":
        if not os.environ.get("GEMINI_API_KEY", "").strip():
            return False, "Hosted Gemini AI is selected, but GEMINI_API_KEY is missing."
        return True, "Hosted Gemini AI is ready."

    try:
        import ollama
    except ImportError:
        return False, "The 'ollama' Python package isn't installed. Run: pip install ollama"

    try:
        from utils import ollama_client
        models_resp = ollama_client.list_models()
    except Exception:
        return False, (
            "Couldn't reach the local Ollama server. Make sure Ollama is installed "
            "and running (open the Ollama app, or run 'ollama serve' in a terminal)."
        )

    # The ollama python client has returned both plain dicts and
    # pydantic-style objects across versions -- handle either.
    raw_models = None
    if isinstance(models_resp, dict):
        raw_models = models_resp.get("models", [])
    else:
        raw_models = getattr(models_resp, "models", []) or []

    available = []
    for m in raw_models:
        if isinstance(m, dict):
            available.append(m.get("model") or m.get("name"))
        else:
            available.append(getattr(m, "model", None) or getattr(m, "name", None))

    model_base = model.split(":")[0]
    has_model = any(a and (a == model or a.split(":")[0] == model_base) for a in available)

    if not has_model:
        return False, (
            f"Ollama is running, but the model '{model}' isn't pulled yet. "
            f"Run: ollama pull {model}"
        )

    return True, "Ollama is ready."


def extract_text(file, progress_callback=None, ocr_model: str = "gemma3:4b"):
    filename = file.name.lower()

    if filename.endswith(".pdf"):
        return read_pdf(file, progress_callback=progress_callback, ocr_model=ocr_model)
    elif filename.endswith(".docx"):
        return read_docx(file)
    elif filename.endswith(".txt"):
        return read_txt(file)
    elif filename.endswith((".jpg", ".jpeg", ".png")):
        return read_image(file, progress_callback=progress_callback, ocr_model=ocr_model)

    return ""
